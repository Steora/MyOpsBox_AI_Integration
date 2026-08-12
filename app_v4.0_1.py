# --- STYLED STREAMLIT APP ENGINE FOR MY OPSBOX ARCHIVER ---
import io
import os
import tempfile
import subprocess
import platform
import PyPDF2
from datetime import datetime, timedelta
import requests
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor


# --- GEMINI MODEL CONFIGURATION ---
PRIMARY_GEMINI_MODEL = "gemini-3.5-flash"
FALLBACK_GEMINI_MODEL = "gemini-3.5-pro-latest"

# --- BRAND PALETTE DEFINITIONS ---
HEX_PRIMARY = "6A398E"                    
COLOR_PRIMARY = RGBColor(106, 57, 142)    
COLOR_SECONDARY = RGBColor(115, 115, 115) 
COLOR_TEXT = RGBColor(60, 60, 60)         
HEX_LIGHT_BG = "F7F8FA"                   

def call_gemini_api(prompt_text, api_key, system_instruction=None, temperature=0.25):
    """Sends a prompt to Google's Gemini REST API and returns the clean generated string."""
    if not api_key:
        raise ValueError("Gemini API key is missing.")

    models_to_try = [PRIMARY_GEMINI_MODEL, FALLBACK_GEMINI_MODEL]
    last_exception = None

    for model in models_to_try:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent?key={api_key}"
        payload = {
            "generationConfig": {"temperature": temperature},
            "contents": [{"role": "user", "parts": [{"text": prompt_text}]}],
        }

        if system_instruction:
            payload["systemInstruction"] = {"parts": [{"text": system_instruction}]}

        try:
            response = requests.post(
                url,
                headers={"Content-Type": "application/json"},
                json=payload,
                timeout=180,
            )

            if response.status_code == 200:
                data = response.json()
                return data["candidates"][0]["content"]["parts"][0]["text"]
            else:
                last_exception = RuntimeError(f"Gemini API ({model}) error {response.status_code}: {response.text}")
        except Exception as err:
            last_exception = err

    raise last_exception or RuntimeError("Failed to communicate with Gemini API endpoints.")


def clean_ai_markdown(text):
    """Removes raw markdown outer code blocks from the AI generated output."""
    if not text:
        return ""
    text = text.strip()
    if text.startswith("```markdown"):
        text = text[11:]
    elif text.startswith("```"):
        text = text[3:]
    if text.endswith("```"):
        text = text[:-3]
    return text.strip()


def extract_pdf_text(uploaded_file):
    """Extracts raw text from an uploaded PDF file for AI processing."""
    if uploaded_file is None:
        return ""
    try:
        reader = PyPDF2.PdfReader(uploaded_file)
        return "\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
    except Exception as e:
        st.error(f"Error extracting PDF: {e}")
        return ""


# --- CONFIGURATION & CONSTANTS ---
FATHOM_API_URL = "[https://api.fathom.ai/external/v1/meetings](https://api.fathom.ai/external/v1/meetings)".strip()

# LOGO & SIGNATURE ASSET PATHS
MY_OPSBOX_LOGO_URL = r"MyOpsBox.svg"  
MY_OPSBOX_LOGO_PNG = r"MyOpsBox.png"
STEORA_LOGO_URL = r"Steora.svg"       
PATTI_SIG_PNG = r"Patricia Zapparolli Signature.png" # Add Signature constant

st.set_page_config(page_title="Discovery Archiver Pro + AI", page_icon=r"Steora-Favicon.png", layout="wide", initial_sidebar_state="collapsed")

# --- BRANDING LAYER: LOGO INTEGRATION ---
logo_col1, logo_col2 = st.columns([1, 1])

with logo_col1:
    try:
        st.image(MY_OPSBOX_LOGO_URL, width=220)
    except Exception:
        st.markdown("### **My OpsBox**")

with logo_col2:
    try:
        if os.path.exists(STEORA_LOGO_URL):
            with open(STEORA_LOGO_URL, "r", encoding="utf-8") as svg_file:
                svg_data = svg_file.read()
            st.markdown(
                f'<div style="display: flex; justify-content: flex-end;"><div style="width: 140px;">{svg_data}</div></div>',
                unsafe_allow_html=True
            )
        else:
            st.markdown("<p style='text-align: right;'><b>Steora Enabled</b></p>", unsafe_allow_html=True)
    except Exception:
        st.markdown("<p style='text-align: right;'><b>Steora Enabled</b></p>", unsafe_allow_html=True)

st.markdown("---")
st.title("🚀 My OpsBox Intelligent Pipeline Engine")
st.markdown("Sync Fathom meetings, run deep-dive strategic analysis, and generate client-ready proposals instantly.")

# --- SIDEBAR CONFIGURATION ---
st.sidebar.header("⚙️ Pipeline Status")

secret_gemini = st.secrets.get("GEMINI_API_KEY") or st.secrets.get("general", {}).get("GEMINI_API_KEY") or os.getenv("GEMINI_API_KEY")
secret_fathom = st.secrets.get("FATHOM_API_KEY") or st.secrets.get("general", {}).get("FATHOM_API_KEY") or os.getenv("FATHOM_API_KEY")

gemini_placeholder = "•••••••••••••••• (Autofetched)" if secret_gemini else "Enter Gemini API Key"
fathom_placeholder = "•••••••••••••••• (Autofetched)" if secret_fathom else "Enter Fathom API Key"

user_gemini = st.sidebar.text_input("Gemini API Key", type="password", placeholder=gemini_placeholder)
user_fathom = st.sidebar.text_input("Fathom API Key", type="password", placeholder=fathom_placeholder)

gemini_credential = user_gemini if user_gemini else secret_gemini
fathom_credential = user_fathom if user_fathom else secret_fathom

st.sidebar.subheader("📅 Sync Parameters")
time_frame = st.sidebar.selectbox("Lookback Window", ["Today", "Past 7 Days", "Past 30 Days", "All Time"], index=3)

if not gemini_credential or not fathom_credential:
    st.info("💡 Please input your Fathom and Gemini API Keys in the sidebar or save them in your Cloud Advanced Settings to activate the automated workspace pipeline.")
    st.stop()

headers = {"X-Api-Key": fathom_credential}

created_after_param = None
if time_frame == "Today":
    created_after_param = (datetime.utcnow() - timedelta(days=1)).strftime("%Y-%m-%dT%H:%M:%SZ")
elif time_frame == "Past 7 Days":
    created_after_param = (datetime.utcnow() - timedelta(days=7)).strftime("%Y-%m-%dT%H:%M:%SZ")
elif time_frame == "Past 30 Days":
    created_after_param = (datetime.utcnow() - timedelta(days=30)).strftime("%Y-%m-%dT%H:%M:%SZ")


# --- XML HELPER FUNCTIONS ---
def set_cell_background(cell, fill_hex):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)


# --- DOCUMENT ENGINE: ASSESSMENT EXPORTER ---
def text_to_docx_buffer(text_content, title_text):
    text_content = clean_ai_markdown(text_content)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    header_table = doc.add_table(rows=1, cols=2)
    header_table.autofit = False
    header_table.rows[0].cells[0].width = Inches(3.8)
    header_table.rows[0].cells[1].width = Inches(2.7)

    left_cell = header_table.rows[0].cells[0]
    left_p = left_cell.paragraphs[0]
    left_p.paragraph_format.space_after = Pt(2)

    if os.path.exists(MY_OPSBOX_LOGO_PNG):
        try:
            left_p.add_run().add_picture(MY_OPSBOX_LOGO_PNG, width=Inches(2.4))
        except Exception:
            brand_run = left_p.add_run("My OpsBox®")
            brand_run.font.name = 'Arial'; brand_run.font.size = Pt(24); brand_run.font.bold = True; brand_run.font.color.rgb = COLOR_PRIMARY
    else:
        brand_run = left_p.add_run("My OpsBox®")
        brand_run.font.name = 'Arial'; brand_run.font.size = Pt(24); brand_run.font.bold = True; brand_run.font.color.rgb = COLOR_PRIMARY

    sub_p = left_cell.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(4)
    sub_run = sub_p.add_run(title_text)
    sub_run.font.name = 'Arial'; sub_run.font.size = Pt(10.5); sub_run.font.italic = True; sub_run.font.color.rgb = COLOR_SECONDARY

    right_cell = header_table.rows[0].cells[1]
    right_p = right_cell.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    tcPr = right_cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="12" w:space="0" w:color="{HEX_PRIMARY}"/></w:tcBorders>')
    tcPr.append(borders)
    set_cell_margins(right_cell, top=60, bottom=60, left=180, right=60)

    for r_text, r_bold, r_color, r_size in [
        ("Prepared by:\n", True, COLOR_SECONDARY, 9.5),
        ("Patti Zapparolli\n", True, COLOR_TEXT, 10),
        ("patti@myopsbox.com\n", False, COLOR_PRIMARY, 9.5),
        ("727-919-7323", False, COLOR_SECONDARY, 9.5)
    ]:
        run = right_p.add_run(r_text)
        run.font.name = 'Arial'; run.font.bold = r_bold; run.font.color.rgb = r_color; run.font.size = Pt(r_size)

    sep_p = doc.add_paragraph()
    sep_p.paragraph_format.space_before = Pt(12)
    sep_p.paragraph_format.space_after = Pt(16)
    sep_p.add_run("―" * 60).font.color.rgb = RGBColor(210, 214, 219)

    lines = text_content.split('\n')
    idx = 0

    while idx < len(lines):
        line = lines[idx].strip()
        if not line or line.startswith("---") or line.startswith("___") or line.startswith(":---"):
            idx += 1
            continue

        if line.startswith("|"):
            table_rows = []
            while idx < len(lines) and lines[idx].strip().startswith("|"):
                row_raw = lines[idx].strip()
                if not (":---" in row_raw or "---:" in row_raw):
                    row_cells = [c.strip() for c in row_raw.split("|")[1:-1]]
                    if row_cells:
                        table_rows.append(row_cells)
                idx += 1

            if table_rows:
                max_cols = max(len(r) for r in table_rows)
                grid_table = doc.add_table(rows=0, cols=max_cols)
                grid_table.style = 'Normal Table'
                set_table_borders(grid_table)

                is_metadata_card = (max_cols == 4)

                for r_num, row_data in enumerate(table_rows):
                    row_cells = grid_table.add_row().cells
                    is_header_row = (r_num == 0 and not is_metadata_card)

                    for c_num, val in enumerate(row_data):
                        if c_num < len(row_cells):
                            cell = row_cells[c_num]
                            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                            cell_p = cell.paragraphs[0]
                            cell_p.paragraph_format.space_after = Pt(0)

                            run = cell_p.add_run(val)
                            run.font.name = 'Arial'

                            if is_header_row:
                                set_cell_background(cell, HEX_PRIMARY)
                                run.font.bold = True; run.font.size = Pt(10); run.font.color.rgb = RGBColor(255, 255, 255)
                            else:
                                run.font.size = Pt(9.5)
                                if is_metadata_card and c_num in [0, 2]:
                                    run.font.bold = True; run.font.color.rgb = COLOR_PRIMARY
                                else:
                                    run.font.color.rgb = COLOR_TEXT

                                if not is_metadata_card and r_num % 2 == 1:
                                    set_cell_background(cell, HEX_LIGHT_BG)

                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                continue

        if line.startswith("## "):
            clean_title = line.replace("## ", "").replace("**", "").strip()
            banner_table = doc.add_table(rows=1, cols=1)
            banner_table.autofit = False
            banner_table.columns[0].width = Inches(6.5)

            cell = banner_table.rows[0].cells[0]
            set_cell_background(cell, HEX_PRIMARY)
            set_cell_margins(cell, top=140, bottom=140, left=120, right=120)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(2); p.paragraph_format.keep_with_next = True

            run = p.add_run(clean_title)
            run.font.name = 'Arial'; run.font.size = Pt(11.5); run.font.bold = True; run.font.color.rgb = RGBColor(255, 255, 255)

            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_after = Pt(4)
            idx += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True

            run = p.add_run(line.replace("### ", "").replace("**", "").strip())
            run.font.name = 'Arial'; run.font.size = Pt(11); run.font.bold = True; run.font.color.rgb = COLOR_PRIMARY
            idx += 1
            continue

        if line.startswith("*") or line.startswith("-") or line.startswith("•"):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(3)

            clean_line = line.lstrip("*•- ").strip()

            if " : " in clean_line or "**" in clean_line:
                clean_line = clean_line.replace("**", "")
                splitter = " : " if " : " in clean_line else ":"
                parts = clean_line.split(splitter, 1)

                r_bold = p.add_run(parts[0] + splitter)
                r_bold.font.name = 'Arial'; r_bold.font.size = Pt(10.5); r_bold.font.bold = True; r_bold.font.color.rgb = COLOR_PRIMARY

                if len(parts) > 1:
                    r_body = p.add_run(parts[1])
                    r_body.font.name = 'Arial'; r_body.font.size = Pt(10.5); r_body.font.color.rgb = COLOR_TEXT
            else:
                run = p.add_run(clean_line)
                run.font.name = 'Arial'; run.font.size = Pt(10.5); run.font.color.rgb = COLOR_TEXT

            idx += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(8); p.paragraph_format.line_spacing = 1.15

        clean_text = line.replace("**", "").strip()

        if ":" in clean_text and len(clean_text.split(":")[0]) < 30 and not clean_text.startswith("http"):
            parts = clean_text.split(":", 1)
            r_bold = p.add_run(parts[0] + ":")
            r_bold.font.name = 'Arial'; r_bold.font.size = Pt(10.5); r_bold.font.bold = True; r_bold.font.color.rgb = COLOR_PRIMARY

            r_body = p.add_run(parts[1])
            r_body.font.name = 'Arial'; r_body.font.size = Pt(10.5); r_body.font.color.rgb = COLOR_TEXT
        else:
            run = p.add_run(clean_text)
            run.font.name = 'Arial'; run.font.size = Pt(10.5); run.font.color.rgb = COLOR_TEXT

        idx += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# --- DOCUMENT ENGINE: CUSTOM PSA EXPORTER ---
def psa_to_docx_buffer(text_content, project_type_footer="PROJECT TYPE: Master Service Agreement | PSA#: TBD"):
    """Dedicated DOCX engine that exactly matches the My OpsBox PSA Template formatting."""
    text_content = clean_ai_markdown(text_content)
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
        # Inject Custom Header with Logo
        header = section.header
        header_p = header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if os.path.exists(MY_OPSBOX_LOGO_PNG):
            try:
                header_p.add_run().add_picture(MY_OPSBOX_LOGO_PNG, width=Inches(2.5))
            except Exception:
                run = header_p.add_run("My OpsBox®")
                run.font.name = 'Arial'; run.font.size = Pt(20); run.font.bold = True; run.font.color.rgb = COLOR_PRIMARY
        
        # Inject Custom Footer
        footer = section.footer
        footer_p = footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f_run = footer_p.add_run(project_type_footer)
        f_run.font.name = 'Arial'; f_run.font.size = Pt(9); f_run.font.color.rgb = COLOR_SECONDARY

    lines = text_content.split('\n')
    idx = 0

    while idx < len(lines):
        line = lines[idx].strip()
        if not line or line.startswith("---") or line.startswith("___") or line.startswith(":---"):
            idx += 1
            continue

        # Intercept the exact Signature Block Marker
        if "[SIGNATURE_BLOCK]" in line:
            doc.add_paragraph().paragraph_format.space_before = Pt(24)
            sig_p = doc.add_paragraph()
            sig_run = sig_p.add_run("Acceptance\nIN WITNESS WHEREOF, the parties have executed this Final Acceptance Section as of the date first written below.")
            sig_run.font.name = 'Arial'; sig_run.font.size = Pt(10)
            
            sig_table = doc.add_table(rows=5, cols=2)
            sig_table.autofit = False
            sig_table.columns[0].width = Inches(3.2)
            sig_table.columns[1].width = Inches(3.2)
            
            # Headers
            sig_table.cell(0,0).text = "Client: ___________________________"
            sig_table.cell(0,1).text = "Provider: My OpsBox"
            
            # Client Signature Line
            sig_table.cell(2,0).text = "_________________________________"
            
            # --- PROVIDER SIGNATURE IMAGE INJECTION ---
            prov_sig_p = sig_table.cell(2,1).paragraphs[0]
            if os.path.exists(PATTI_SIG_PNG):
                prov_sig_p.add_run().add_picture(PATTI_SIG_PNG, width=Inches(1.8))
            else:
                prov_sig_p.add_run("_________________________________")
            
            # Name and Title Lines
            sig_table.cell(3,0).text = "Client Representative Name and Title"
            sig_table.cell(3,1).text = "Patti Zapparolli, Founder\nPROVIDER Representative Name and Title"
            
            # Date Lines
            sig_table.cell(4,0).text = "___________________________________\nDate Signed"
            sig_table.cell(4,1).text = f"{datetime.now().strftime('%B %d, %Y')}\nDate Signed"
            
            # Apply styling to the entire signature table
            for row in sig_table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = 'Arial'
                            run.font.size = Pt(10)
            idx += 1
            continue

        if line.startswith("## "):
            clean_title = line.replace("## ", "").replace("**", "").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18); p.paragraph_format.space_after = Pt(6); p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_title)
            run.font.name = 'Arial'; run.font.size = Pt(16); run.font.bold = True; run.font.color.rgb = COLOR_PRIMARY
            
            # Add thin separator line under header
            sep_p = doc.add_paragraph()
            sep_p.paragraph_format.space_before = Pt(0); sep_p.paragraph_format.space_after = Pt(8)
            sep_p.add_run("―" * 60).font.color.rgb = RGBColor(210, 214, 219)
            idx += 1
            continue

        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4); p.paragraph_format.keep_with_next = True
            run = p.add_run(line.replace("### ", "").replace("**", "").strip())
            run.font.name = 'Arial'; run.font.size = Pt(12); run.font.bold = True; run.font.color.rgb = COLOR_TEXT
            idx += 1
            continue

        if line.startswith("*") or line.startswith("-") or line.startswith("•"):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(3)
            clean_line = line.lstrip("*•- ").strip()

            if " : " in clean_line or "**" in clean_line:
                clean_line = clean_line.replace("**", "")
                splitter = " : " if " : " in clean_line else ":"
                parts = clean_line.split(splitter, 1)

                r_bold = p.add_run(parts[0] + splitter)
                r_bold.font.name = 'Arial'; r_bold.font.size = Pt(10.5); r_bold.font.bold = True; r_bold.font.color.rgb = COLOR_TEXT
                if len(parts) > 1:
                    r_body = p.add_run(parts[1])
                    r_body.font.name = 'Arial'; r_body.font.size = Pt(10.5)
            else:
                run = p.add_run(clean_line)
                run.font.name = 'Arial'; run.font.size = Pt(10.5)
            idx += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(8); p.paragraph_format.line_spacing = 1.15
        clean_text = line.replace("**", "").strip()

        if ":" in clean_text and len(clean_text.split(":")[0]) < 30 and not clean_text.startswith("http"):
            parts = clean_text.split(":", 1)
            r_bold = p.add_run(parts[0] + ":")
            r_bold.font.name = 'Arial'; r_bold.font.size = Pt(10.5); r_bold.font.bold = True
            r_body = p.add_run(parts[1])
            r_body.font.name = 'Arial'; r_body.font.size = Pt(10.5)
        else:
            run = p.add_run(clean_text)
            run.font.name = 'Arial'; run.font.size = Pt(10.5)

        idx += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def docx_to_pdf_buffer(docx_buffer):
    """Converts a highly-styled DOCX memory buffer into a PDF using system layout engines."""
    with tempfile.TemporaryDirectory() as tmpdir:
        docx_path = os.path.join(tmpdir, "temp.docx")
        pdf_path = os.path.join(tmpdir, "temp.pdf")

        with open(docx_path, "wb") as f:
            f.write(docx_buffer.getvalue())

        if platform.system() == "Windows":
            try:
                from docx2pdf import convert
                import pythoncom 
                pythoncom.CoInitialize()
                convert(docx_path, pdf_path)
                pythoncom.CoUninitialize()
            except ImportError:
                st.error("Missing local package. Run: pip install docx2pdf")
                return None
            except Exception as win_err:
                st.error(f"Local Word Conversion Error: {win_err}")
                return None
        else:
            try:
                subprocess.run(
                    ["libreoffice", "--headless", "--convert-to", "pdf", docx_path, "--outdir", tmpdir],
                    check=True,
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL
                )
            except FileNotFoundError:
                st.error("LibreOffice missing. Ensure packages.txt exists with 'libreoffice'.")
                return None

        with open(pdf_path, "rb") as f:
            return io.BytesIO(f.read())


# --- CENTRAL PIPELINE MECHANICS ---
if st.button("🔄 Sync Active Discovery Call Stream", type="primary"):
    with st.spinner("Connecting to Fathom Cloud Gateways..."):
        try:
            clean_endpoint = FATHOM_API_URL.split("://")[-1].replace(")", "").replace("]", "").strip()
            endpoint_url = f"https://{clean_endpoint}"
            
            params = {"include_transcript": "true", "calendar_invitees_domains_type": "all"}
            if created_after_param:
                params["created_after"] = created_after_param

            response = requests.get(endpoint_url, headers=headers, params=params, timeout=30)
            if response.status_code == 200:
                st.session_state["meetings_list"] = response.json().get("items", [])
                st.success(f"Synced {len(st.session_state['meetings_list'])} calls from Fathom stream.")
            else:
                st.error(f"Fathom error authentication drop: Code {response.status_code}")
        except Exception as e:
            st.error(f"Pipeline failure: {e}")

if "meetings_list" in st.session_state and st.session_state["meetings_list"]:
    meetings = st.session_state["meetings_list"]
    options = [(f"{m.get('created_at', 'TBD')[:10]} — {m.get('title', 'Call')}", idx) for idx, m in enumerate(meetings)]

    selected_call = st.selectbox("Select Call Stream for Analysis:", options=options, format_func=lambda x: x[0])

    if selected_call is not None:
        target = meetings[selected_call[1]]
        rec_id = target.get("recording_id")

        raw_duration = target.get('duration') or target.get('duration_seconds') or target.get('length') or 0
        try:
            duration_seconds = int(raw_duration)
        except (ValueError, TypeError):
            duration_seconds = 0
        duration_minutes = duration_seconds // 60

        if duration_minutes == 0 and "transcript" in target and isinstance(target["transcript"], list):
            try:
                last_turn = target["transcript"][-1]
                ts = last_turn.get("timestamp", "0:00").split(":")
                if len(ts) == 2:
                    duration_minutes = int(ts[0])
                elif len(ts) == 3:
                    duration_minutes = (int(ts[0]) * 60) + int(ts[1])
            except Exception:
                duration_minutes = 15

        col_m1, col_m2 = st.columns(2)
        current_call_title = target.get("title", "Unknown Call")
        col_m1.metric("Selected Client Call", current_call_title)
        col_m2.metric("Duration Pool", f"{duration_minutes} minutes")

        raw_transcript_block = ""
        if rec_id:
            with st.spinner("Downloading raw call data elements..."):
                clean_rec_path = f"api.fathom.ai/external/v1/recordings/{rec_id}/transcript".strip()
                transcript_url = f"https://{clean_rec_path}"
                t_resp = requests.get(transcript_url, headers=headers, timeout=30)
                
                if t_resp.status_code == 200:
                    turns = t_resp.json().get("transcript", [])
                    raw_transcript_block = "\n".join([f"[{t.get('timestamp', '00:00')}] {t.get('speaker', {}).get('display_name', 'Speaker')}: {t.get('text', '')}" for t in turns])
                else:
                    raw_transcript_block = str(target.get("transcript", ""))

        if not raw_transcript_block.strip() or len(raw_transcript_block) < 20:
            if isinstance(target.get("transcript"), list):
                raw_transcript_block = "\n".join([f"{t.get('speaker', 'Speaker')}: {t.get('text', '')}" for t in target["transcript"]])
            else:
                raw_transcript_block = str(target.get("transcript", "No transcript payload present."))

        st.markdown("---")
        st.subheader("🧠 Phase 1: Growth Drag Audit")
        st.markdown("Pass raw transcript blocks through your Mastermind prompt structure directly to generate the initial Assessment.")

        if st.button("🔥 Run AI Assessment Analysis", type="secondary"):
            if not raw_transcript_block or len(raw_transcript_block) < 30:
                st.error("No valid transcript data found for this call window. Verify Fathom processing state.")
            else:
                with st.spinner(f"Processing deep analysis matrix for {current_call_title}..."):
                    try:
                        master_prompt = f"""
You are an expert Fractional Chief Operating Officer (Fractional COO) acting on behalf of My OpsBox. 
Your objective is to thoroughly analyze the attached raw Fathom call transcript and generate a comprehensive Operational Assessment document that matches the strict tone, structure, and depth shown in your brand's note sheets.

### CORE OPERATIONAL SOLUTIONS KNOWLEDGE:
- Leverage the 'My OpsBox Managed Services Guide' positioning matrix. 
- Ensure recommendations cross-reference solutions to explicit tiers.

### CRITICAL OUTPUT FORMATTING INSTRUCTIONS FOR THE DOCX PROCESSING ENGINE:
1. CLIENT CARD META BLOCK:
| Business Name | [Insert] | Assessment Date | {datetime.now().strftime('%m/%d/%y')} |
| Website | [Insert] | Email | [Insert] |
| Business Type | [Insert] | Business Industry | [Insert] |
| Number of full-time employees | [ headcount ] | Annual revenue range | [Estimated] |

2. PRIMARY HEADINGS: (Prevent orphan headings by starting a header on a new page if its content cannot start/fit below it.)
## 1. Executive Summary & Business Context
## 2. Key Operational Challenges & Gaps
## 3. Strategic 12-Month Objectives
## 4. Proposed Service Package Strategy
## 5. Structured 30-60-90 Day Transformation Timeline Plan
## 6. Scorecard Matrix Projections

3. STRICT NO ASCII ART OR FLOWCHARTS:
Never use ASCII block shapes or lines (e.g., ┌───┐). Represent all strategic objectives, matrices, and hierarchies strictly as standard Markdown tables (`| Col | Col |`) or standard bulleted lists.

### LIVE CALL TRANSCRIPT TEXT TO PROCESS AND ANALYZE:
{raw_transcript_block}
"""
                        system_instruction = "You are a world-class Fractional COO and master systems operations analyst specializing in corporate scaling, workflow architecture, and business metrics mapping."
                        ai_response_text = call_gemini_api(master_prompt, gemini_credential, system_instruction=system_instruction)

                        st.session_state["compiled_analysis"] = clean_ai_markdown(ai_response_text)
                        st.success("AI Strategic Assessment Generated Flawlessly From Live Transcript!")
                        st.rerun()

                    except Exception as ai_err:
                        st.error(f"Gemini Core Engine Exception: {ai_err}")

        # === PHASE 1 EXPORTS ===
        if "compiled_analysis" in st.session_state:
            st.markdown("### 📋 Generated Strategic Analysis")
            st.markdown(st.session_state["compiled_analysis"])
            st.markdown("### 📥 Assessment Export Options")
            safe_name = "".join([c if c.isalnum() else "_" for c in target.get('title', 'Analysis')])

            if "export_docx_buf" not in st.session_state or st.session_state.get("last_call_title") != current_call_title:
                st.session_state["export_docx_buf"] = text_to_docx_buffer(st.session_state["compiled_analysis"], "Operational Assessment & Strategic Analysis")
                with st.spinner("Rendering exact-match PDF from Word template..."):
                    st.session_state["export_pdf_buf"] = docx_to_pdf_buffer(st.session_state["export_docx_buf"])
                st.session_state["last_call_title"] = current_call_title

            col_ex1, col_ex2 = st.columns(2)
            with col_ex1:
                st.download_button(label="💾 Save Assessment (Word)", data=st.session_state["export_docx_buf"], file_name=f"AI_Assessment_{safe_name}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            with col_ex2:
                if st.session_state["export_pdf_buf"]:
                    st.download_button(label="💾 Save Assessment (PDF)", data=st.session_state["export_pdf_buf"], file_name=f"AI_Assessment_{safe_name}.pdf", mime="application/pdf")

            # === PHASE 2 PSA GENERATOR ===
            st.markdown("---")
            st.subheader("🤝 Phase 2: Proposal & Service Agreement (PSA) Generator")
            st.markdown("Combine your generated Assessment with your Master Templates to create a fully customized PSA ready for signature.")

            st.markdown("#### 📁 Resources Area")
            col_r1, col_r2 = st.columns(2)
            with col_r1:
                managed_services_file = st.file_uploader("Upload Managed Services Guide", type=["pdf"], key="msg_file")
            with col_r2:
                psa_template_file = st.file_uploader("Upload Sample PSA Template", type=["pdf"], key="psa_file")

            st.markdown("#### ⚙️ Business Rules & Pricing")
            col_p1, col_p2 = st.columns(2)
            with col_p1:
                contract_length = st.text_input("Minimum Contract Length", value="12 months")
            with col_p2:
                pricing_structure = st.text_area("Pricing / Rate Structure Per Month", placeholder="Enter the quoted monthly rate or pricing structure...")

            if st.button("📑 Generate Final PSA Document", type="primary"):
                if not managed_services_file or not psa_template_file:
                    st.warning("Please upload both the Managed Services Guide and the Sample PSA Template PDFs in the Resources Area.")
                elif not pricing_structure.strip():
                    st.warning("Please provide the pricing rate structure to populate the agreement.")
                else:
                    with st.spinner("Extracting template rules and generating highly customized PSA..."):
                        try:
                            msg_text = extract_pdf_text(managed_services_file)
                            psa_text = extract_pdf_text(psa_template_file)

                            psa_prompt = f"""
You are an expert Fractional Chief Operating Officer (Fractional COO) and elite sales closer acting on behalf of My OpsBox.
Your objective is to generate a highly professional, combined Proposal and Service Agreement (PSA) designed to close the full package with the prospect.

For PSA Document - Replicate the attached document’s exact layout, fonts, and inline CSS styling as a 1:1 template. Maintain complete visual fidelity while replacing the text with dynamic data fetched from transcript and assessment analysis.

### SOURCE MATERIALS & CONTEXT:
1. **Assessment Analysis & Discovery Notes:**
{st.session_state["compiled_analysis"]}

2. **Managed Services Guide (for custom Scope of Work mapping):**
{msg_text}

3. **Sample PSA Template Context:**
{psa_text}

4. **Pricing & Rate Structure:**
{pricing_structure}

5. **Specific Business Rules to Enforce:**
- Minimum Contract Length: {contract_length}
- The ultimate goal of this proposal is to close the full package.

### INSTRUCTIONS FOR DOCUMENT GENERATION (STRICT):
Output pure markdown. You must strictly mirror the section headers exactly as shown below so the DOCX engine applies the correct purple corporate branding. Never use ASCII art tables, boxes, or lines.

**Prepared For:** [Insert Client Name, Company Name]
**My PreOpsBox:** [Insert Proposed Role/Package]
**Date Prepared:** {datetime.now().strftime('%B %d, %Y')}

This Proposal and Service Agreement for the [Insert Service] (the "Agreement") will be effective from [Insert Start Date] ("Effective Date") and continue until [Insert End Date], unless terminated by this Agreement's terms and conditions.

**BETWEEN:**
[Insert Company Name] (The "CLIENT"), a company organized and existing under the laws of the State of [Insert State], with its head office located at: [Insert Address]

**AND:**
My OpsBox (The "PROVIDER"), a company organized and existing under the laws of Nevada, with its head office located at: 1900 N Torrey Pines Dr., Unit 215, Las Vegas, NV 89108.

**Service Length:** {contract_length}
**Monthly Rate:** {pricing_structure}

**Scope Transparency:**
- Rate covers only items listed in Scope of Service.
- Additional projects outside of the scope will require additional proposals and costs.
- Rate includes $50 monthly phone fee.

## Scope of Work
[Generate custom scope mapped from the Managed Services Guide based on client needs]

## Value
[Generate custom ROI and value proposition based on the Assessment]

## Time and Payment
[Include standard payment terms from template, adjusting for pricing structure]

## Standard Operating Procedures
[Include standard operating procedures from the template]

## Legal Notices
[Include all standard legal clauses from the template]

## Entire Agreement
[Include the final Entire Agreement and Severability clauses]

[SIGNATURE_BLOCK]
"""
                            system_instruction = "You are a master fractional executive and sales closer. Output pure Markdown documents only."
                            psa_response_text = call_gemini_api(psa_prompt, gemini_credential, system_instruction=system_instruction)

                            st.session_state["compiled_psa"] = clean_ai_markdown(psa_response_text)
                            st.success("Proposal & Service Agreement Generated Successfully!")
                            st.rerun()

                        except Exception as e:
                            st.error(f"PSA Generation Error: {e}")

            # === PHASE 2 EXPORTS ===
            if "compiled_psa" in st.session_state:
                st.markdown("### 📋 Generated Proposal & Service Agreement")
                st.markdown(st.session_state["compiled_psa"])
                st.markdown("### 📥 PSA Export Options")
                
                project_footer = f"PROJECT TYPE: Proposal & Service Agreement | PSA#: {safe_name.upper()}"
                
                if "export_psa_docx" not in st.session_state or st.session_state.get("last_psa_title") != current_call_title:
                    st.session_state["export_psa_docx"] = psa_to_docx_buffer(st.session_state["compiled_psa"], project_footer)
                    with st.spinner("Rendering exact-match PDF for PSA..."):
                        st.session_state["export_psa_pdf"] = docx_to_pdf_buffer(st.session_state["export_psa_docx"])
                    st.session_state["last_psa_title"] = current_call_title

                col_psa1, col_psa2 = st.columns(2)
                with col_psa1:
                    st.download_button(label="💾 Save PSA (Word)", data=st.session_state["export_psa_docx"], file_name=f"PSA_{safe_name}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                with col_psa2:
                    if st.session_state["export_psa_pdf"]:
                        st.download_button(label=" 💾 Save PSA (PDF)", data=st.session_state["export_psa_pdf"], file_name=f"PSA_{safe_name}.pdf", mime="application/pdf")